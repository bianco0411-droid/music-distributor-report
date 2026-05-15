"""
只做一件事：重新计算所有 h3 标题的百分比（不合并 OC/UGC），然后替换 HTML 中对应的标题行。
表格数据（DATA 对象）这次不动，下次再处理。
"""
import openpyxl, re

EUR_TO_USD = 1.085

# ========== 1. 从 Believe xlsx 算各 Platform 总收入 ==========
wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', data_only=True)
ws = wb.active
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
plat_idx  = headers.index('Platform')
gross_idx = headers.index('Gross Revenue')

believe_total_by_plat = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    p = row[plat_idx]
    g = row[gross_idx]
    if p and g is not None:
        believe_total_by_plat[p] = believe_total_by_plat.get(p, 0.0) + float(g) * EUR_TO_USD

believe_all = sum(believe_total_by_plat.values())
print(f'Believe 总收入: {believe_all:,.2f}')

# ========== 2. 从东西 CSV 算各 Store 总收入 ==========
import csv
dongxi_total_by_store = {}
dongxi_all = 0.0
for f in ['/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202601.csv',
          '/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202602.csv',
          '/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202603.csv']:
    try:
        with open(f, encoding='utf-8-sig') as fh:
            for row in csv.DictReader(fh):
                s = row.get('Store','')
                try:
                    r = float(row.get('AmountPaidByStore',0))
                except: r = 0
                dongxi_total_by_store[s] = dongxi_total_by_store.get(s,0.0) + r
                dongxi_all += r
    except: pass

print(f'东西总收入:   {dongxi_all:,.2f}')
total_all = believe_all + dongxi_all
print(f'合计:         {total_all:,.2f}')
print()

# ========== 3. 以实际名称为准，计算各板块百分比 ==========
# 板块定义：（显示名，emoji标题关键词，Believe Platform，东西 Store）
sections = [
    ('Spotify',      '🟢 Spotify',              'Spotify',              'Spotify'),
    ('Apple Music',  '🍎 Apple Music',          'Apple Music',          'Apple Music'),
    ('YT OC vs AT', '🎬 YouTube Official Content vs Art Tracks', 'YouTube Official Content', 'YouTube Art Tracks'),
    ('YT UGC vs ACID', '🎬 YouTube Content ID', 'YouTube UGC',       'YouTube Audio Content ID'),
    ('YT Audio Tier','▶️ YouTube Audio Tier',   'YouTube Audio Tier',   'YouTube Audio Tier'),
    ('YT Music Vid', '🎵 YouTube Music Video',  'YouTube Music Video',  'YouTube Music Video'),
    ('YT Shorts',   '🩳 YouTube Shorts',       'YouTube Shorts',       'YouTubeShorts'),
]

print(f'{"板块":35s} {"B(USD)":>14s} {"D(USD)":>14s} {"合计":>14s} {"占总收入%":>10s} {"东西%":>8s} {"Believe%":>10s}')
print('-' * 120)

results = {}
for name, emoji_h3, b_plat, d_store in sections:
    b = believe_total_by_plat.get(b_plat, 0.0)
    d = dongxi_total_by_store.get(d_store, 0.0)
    combined = b + d
    pct_total = combined / total_all * 100 if total_all else 0
    pct_d = d / combined * 100 if combined else 0
    pct_b = b / combined * 100 if combined else 0
    print(f'{name:35s} {b:>14,.2f} {d:>14,.2f} {combined:>14,.2f} {pct_total:>9.1f}% {pct_d:>7.1f}% {pct_b:>9.1f}%')
    results[emoji_h3] = (pct_total, pct_d, pct_b)

# ========== 4. 替换 HTML 中的 h3 标题 ==========
with open('/Users/olivia/WorkBuddy/2026-05-13-task-4/index.html', 'r', encoding='utf-8') as f:
    lines = f.readlines()

new_lines = []
for line in lines:
    matched = False
    for emoji_h3, (pt, pd, pb) in results.items():
        # 匹配包含 emoji_h3 和 "每1,000" 的行
        if emoji_h3 in line and '每1,000' in line:
            # 找行中 emoji_h3 的位置，重建整行
            # 格式：<h3 ...>emoji_h3 — 每1,000 xxx（占总收入 ...%</h3>
            # 我们要替换掉（占总收入 ...%）这段
            suffix = 'streams' if 'streams' in line else 'shorts'
            new_text = f'{emoji_h3} — 每1,000 {suffix}（占总收入 {pt:.1f}%，其中东西占{pd:.1f}%，Believe占{pb:.1f}%）'
            # 替换 > 和 </h3> 之间的内容
            import re
            new_line = re.sub(r'>[^<]+</h3>', f'>{new_text}</h3>', line)
            new_lines.append(new_line)
            matched = True
            break
    if not matched:
        new_lines.append(line)

with open('/Users/olivia/WorkBuddy/2026-05-13-task-4/index.html', 'w', encoding='utf-8') as f:
    f.writelines(new_lines)

print()
print('HTML h3 标题已更新！')
print('注意：')
print('- YouTube Official Content 和 YouTube UGC 已分开计算')
print('- "YouTube Content ID" 板块现在代表 UGC vs Audio Content ID')
print('- "YouTube Official Content vs Art Tracks" 板块现在只用 Official Content（不合并）')
print('- 表格数据（DATA 对象）本次未改动，如需改动请告知。')

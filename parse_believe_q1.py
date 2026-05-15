import openpyxl
import json
from collections import defaultdict

# 加载 Excel（read_only 模式处理大文件）
print("正在加载 Excel 文件...")
wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', read_only=True)
ws = wb.active
print(f"Sheet: {ws.title}, 总行数: {ws.max_row}")

# 读取表头，找到各列索引
headers = None
for row in ws.iter_rows(min_row=1, max_row=1):
    headers = [cell.value for cell in row]
print("表头:", headers)

# 找到关键列的索引
idx_platform = headers.index('Platform')
idx_country = headers.index('Country / Region')
idx_quantity = headers.index('Quantity')
idx_gross = headers.index('Gross Revenue')

# 地区名 → 代码映射
COUNTRY_MAP = {
    'united states': 'US', 'usa': 'US', 'us': 'US',
    'united kingdom': 'GB', 'uk': 'GB', 'great britain': 'GB',
    'japan': 'JP',
    'korea': 'KR', 'korea, republic of': 'KR', 'korea republic of': 'KR', 'republic of korea': 'KR', 'south korea': 'KR',
    'taiwan': 'TW', 'taiwan, province of china': 'TW',
    'thailand': 'TH',
    'indonesia': 'ID',
    'viet nam': 'VN', 'vietnam': 'VN',
    'hong kong': 'HK', 'hong kong sar': 'HK',
    'singapore': 'SG',
    'malaysia': 'MY',
    'china': 'CN', 'china mainland': 'CN', 'mainland china': 'CN',
    'germany': 'DE', 'deutschland': 'DE',
    'france': 'FR',
    'australia': 'AU',
    'canada': 'CA',
    'brazil': 'BR',
    'mexico': 'MX',
}

# Platform 映射：Believe Platform 名 → DATA 中的 key 前缀
# bSp, bAp, bYm, bYv, bYc, bSh
PLATFORM_MAP = {
    'Spotify': 'Sp',
    'Apple Music': 'Ap',
    'Apple iTunes': 'Ap',        # iTunes 归入 Apple
    'YouTube Official Content': 'Yc',   # Content ID
    'YouTube UGC': 'Yc',               # UGC 也归入 Content ID
    'YouTube Audio Tier': 'Ym',        # Audio Tier
    'YouTube Music Video': 'Yv',       # Music Video
    'YouTube Shorts': 'Sh',            # Shorts
    'TikTok': 'Tk',
    'TikTok TV': 'Tk',
    'Facebook': 'Fb',
    'Instagram': 'Fb',
}

# 汇总：{region: {platform_key: [total_gross_eur, total_qty]}}
agg = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))

row_count = 0
korea_rows = 0

for row in ws.iter_rows(min_row=2, values_only=True):
    row_count += 1
    if row_count % 100000 == 0:
        print(f"  已处理 {row_count} 行...")
    
    platform = row[idx_platform]
    country = row[idx_country]
    qty = row[idx_quantity]
    gross = row[idx_gross]
    
    # 跳过空行
    if platform is None or country is None:
        continue
    
    # 地区映射
    country_key = str(country).strip().lower()
    region = COUNTRY_MAP.get(country_key)
    
    # 统计韩国行数
    if 'korea' in country_key:
        korea_rows += 1
    
    if region is None:
        continue
    
    # Platform 映射
    platform_key = PLATFORM_MAP.get(str(platform).strip())
    if platform_key is None:
        continue
    
    # 累加
    try:
        q = float(qty) if qty is not None else 0.0
        g = float(gross) if gross is not None else 0.0
        if q > 0:  # 只统计有播放量的
            agg[region][platform_key][0] += g  # gross revenue EUR
            agg[region][platform_key][1] += q  # quantity
    except (TypeError, ValueError):
        pass

wb.close()

print(f"\n总处理行数: {row_count}")
print(f"韩国相关行数: {korea_rows}")

# 计算单价并输出
print("\n=== Believe 各地区各平台单价（USD / 1000 streams）===")
print("汇率: 1 EUR = 1.085 USD")
print()

RESULTS = {}
for region in sorted(agg.keys()):
    RESULTS[region] = {}
    print(f"[{region}]")
    for pk in sorted(agg[region].keys()):
        total_gross, total_qty = agg[region][pk]
        if total_qty > 0:
            # 加权平均单价 = sum(Gross EUR) × 1.085 / sum(Qty) × 1000
            unit_price = (total_gross * 1.085) / total_qty * 1000
            RESULTS[region][pk] = round(unit_price, 4)
            print(f"  b{pk}: {unit_price:.4f}")
        else:
            RESULTS[region][pk] = None
            print(f"  b{pk}: N/A (no data)")
    print()

# 输出 JSON 供复制
print("=== JSON 格式（用于更新 DATA 对象）===")
print(json.dumps(RESULTS, indent=2, ensure_ascii=False))

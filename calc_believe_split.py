import openpyxl
from collections import defaultdict

EUR_TO_USD = 1.085

wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', data_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

plat_idx = headers.index('Platform')
country_idx = headers.index('Country / Region')
qty_idx = headers.index('Quantity')
gross_idx = headers.index('Gross Revenue')

COUNTRY_MAP = {
    'united states': 'US',
    'united kingdom': 'GB',
    'japan': 'JP',
    'korea, republic of': 'KR',
    'taiwan, province of china': 'TW',
    'thailand': 'TH',
    'indonesia': 'ID',
    'viet nam': 'VN',
    'hong kong': 'HK',
    'singapore': 'SG',
    'malaysia': 'MY',
    'china': 'CN',
}

# data[region][platform] = {rev, qty}
data = defaultdict(lambda: defaultdict(lambda: {'rev': 0.0, 'qty': 0.0}))

for row in ws.iter_rows(min_row=2, values_only=True):
    plat = row[plat_idx]
    country = row[country_idx]
    qty = row[qty_idx]
    gross = row[gross_idx]
    if not plat or not country or qty is None or gross is None:
        continue
    region = COUNTRY_MAP.get(str(country).strip().lower())
    if not region:
        continue
    try:
        qty_f = float(qty)
        rev_f = float(gross) * EUR_TO_USD
    except (ValueError, TypeError):
        continue
    data[region][plat]['rev'] += rev_f
    data[region][plat]['qty'] += qty_f

def unit_price(rev, qty):
    return rev / qty * 1000 if qty > 0 else None

def fmt(v):
    return f'{v:>20.4f}' if v is not None else f'{"N/A":>20s}'

regions = ['US', 'GB', 'JP', 'KR', 'TW', 'TH', 'ID', 'VN', 'HK', 'SG', 'MY']

print('Believe 分地区分Platform单价（$/1000）：')
print(f'{"Region":8s} {"Official Content":>20s} {"UGC":>20s} {"Audio Tier":>20s} {"Music Video":>20s} {"Shorts":>20s}')
print('-' * 120)

for r in regions:
    rd = data[r]
    oc_up = unit_price(rd['YouTube Official Content']['rev'], rd['YouTube Official Content']['qty'])
    ugc_up = unit_price(rd['YouTube UGC']['rev'], rd['YouTube UGC']['qty'])
    at_up = unit_price(rd['YouTube Audio Tier']['rev'], rd['YouTube Audio Tier']['qty'])
    mv_up = unit_price(rd['YouTube Music Video']['rev'], rd['YouTube Music Video']['qty'])
    sh_up = unit_price(rd['YouTube Shorts']['rev'], rd['YouTube Shorts']['qty'])
    print(f'{r:8s} {fmt(oc_up)} {fmt(ugc_up)} {fmt(at_up)} {fmt(mv_up)} {fmt(sh_up)}')

# Totals
print()
print('=== 合计（所有地区，加权平均）===')
for plat in ['YouTube Official Content', 'YouTube UGC', 'YouTube Audio Tier', 'YouTube Music Video', 'YouTube Shorts']:
    total_rev = sum(data[r][plat]['rev'] for r in regions)
    total_qty = sum(data[r][plat]['qty'] for r in regions)
    up = unit_price(total_rev, total_qty)
    print(f'  {plat}: {up:.4f}/1000' if up else f'  {plat}: N/A')

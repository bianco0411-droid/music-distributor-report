import openpyxl
import csv
from collections import defaultdict

EUR_TO_USD = 1.085

# ========== 1. Believe 数据（从 xlsx，按实际 Platform 名称）==========
wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', data_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
platform_idx = headers.index('Platform')
gross_idx = headers.index('Gross Revenue')

believe_rev = defaultdict(float)
for row in ws.iter_rows(min_row=2, values_only=True):
    p = row[platform_idx]
    g = row[gross_idx]
    if p and g is not None:
        try:
            believe_rev[p] += float(g) * EUR_TO_USD
        except: pass

# ========== 2. 东西数据（从 CSV，按实际 Store 名称）==========
dongxi_rev = defaultdict(float)
for f in ['/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202601.csv',
          '/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202602.csv',
          '/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_202603.csv']:
    with open(f, encoding='utf-8-sig') as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            s = row.get('Store', '')
            try:
                r = float(row.get('AmountPaidByStore', 0))
            except: r = 0
            dongxi_rev[s] += r

# ========== 3. 以实际名称对应，不合并 ==========
# 用户要求：以报表实际 store/platform 为主，UGC 和 Official Content 不合并
# Believe YouTube Official Content ↔ 东西 YouTube Art Tracks
# Believe YouTube UGC ↔ 东西 YouTube Audio Content ID
# Believe YouTube Audio Tier ↔ 东西 YouTube Audio Tier
# Believe YouTube Music Video ↔ 东西 YouTube Music Video
# Believe YouTube Shorts ↔ 东西 YouTubeShorts

comparisons = [
    ('Spotify',            'Spotify',               'Spotify'),
    ('Apple Music',        'Apple Music',            'Apple Music'),
    ('YouTube Official Content', 'YouTube Official Content', 'YouTube Art Tracks'),
    ('YouTube UGC',       'YouTube UGC',            'YouTube Audio Content ID'),
    ('YouTube Audio Tier', 'YouTube Audio Tier',     'YouTube Audio Tier'),
    ('YouTube Music Video','YouTube Music Video',     'YouTube Music Video'),
    ('YouTube Shorts',    'YouTube Shorts',         'YouTubeShorts'),
    ('Facebook / Instagram','Facebook / Instagram',   'Facebook'),
    ('TikTok',            'TikTok',                'TikTok'),
]

# 计算总收入（用于占总收入百分比）
all_believe = sum(believe_rev.values())
all_dongxi = sum(dongxi_rev.values())
total_all = all_believe + all_dongxi

print(f'Believe 总收入: {all_believe:,.2f} USD')
print(f'东西总收入:   {all_dongxi:,.2f} USD')
print(f'合计:         {total_all:,.2f} USD')
print()

print('=== 各平台对比（不合并 UGC/OC）===')
print(f'{"对比板块":40s} {"Believe(USD)":>16s} {"东西(USD)":>16s} {"合计":>16s} {"占总收入%":>10s} {"东西%":>8s} {"Believe%":>10s}')
print('-' * 120)

results = {}
for name, b_plat, d_store in comparisons:
    b = believe_rev.get(b_plat, 0.0)
    d = dongxi_rev.get(d_store, 0.0)
    combined = b + d
    pct_total = combined / total_all * 100 if total_all else 0
    pct_d = d / combined * 100 if combined else 0
    pct_b = b / combined * 100 if combined else 0
    print(f'{name:40s} {b:>16,.2f} {d:>16,.2f} {combined:>16,.2f} {pct_total:>9.1f}% {pct_d:>7.1f}% {pct_b:>9.1f}%')
    results[name] = {
        'b_plat': b_plat,
        'd_store': d_store,
        'b': b,
        'd': d,
        'combined': combined,
        'pct_total': pct_total,
        'pct_d': pct_d,
        'pct_b': pct_b,
    }

print()
print('注意：YouTube Official Content 和 YouTube UGC 已分开对比，未合并。')

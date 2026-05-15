import openpyxl
import json
from collections import defaultdict

EUR_TO_USD = 1.085

# === Load Believe data from xlsx ===
wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', data_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
platform_idx = headers.index('Platform')
gross_idx = headers.index('Gross Revenue')

platform_revenue = defaultdict(float)
for row in ws.iter_rows(min_row=2, values_only=True):
    platform = row[platform_idx]
    gross = row[gross_idx]
    if platform and gross is not None:
        try:
            platform_revenue[platform] += float(gross)
        except (ValueError, TypeError):
            pass

# Map Believe platforms to DATA keys
believe_mapping = {
    'Spotify': 'Sp',
    'Apple Music': 'Ap',
    'YouTube Official Content': 'Yc',
    'YouTube UGC': 'Yc',
    'YouTube Audio Tier': 'Ym',
    'YouTube Music Video': 'Yv',
    'YouTube Shorts': 'Sh',
    'Facebook / Instagram': 'Fb',
    'TikTok': 'Tk',
}

b_data = {}
for platform, key in believe_mapping.items():
    rev = platform_revenue.get(platform, 0.0) * EUR_TO_USD
    b_data[key] = b_data.get(key, 0.0) + rev

# === Load 东西 data ===
with open('/Users/olivia/WorkBuddy/2026-05-13-task-4/dongxi_new_analysis.json') as f:
    dx = json.load(f)

dx_summary = {item['platform']: item['revenue'] for item in dx['platform_summary']}

# Map 东西 platforms to DATA keys
dx_mapping = {
    'Spotify': 'Sp',
    'Apple Music': 'Ap',
    'YouTube Content ID': 'Yc',
    'YouTube Music': 'Ym',
    'YouTube Music Video': 'Yv',
    'YouTube Shorts': 'Sh',
    'Facebook': 'Fb',
    'TikTok': 'Tk',
}

d_data = {}
for platform, key in dx_mapping.items():
    d_data[key] = dx_summary.get(platform, 0.0)

# === Calculate percentages ===
total_rev = sum(b_data.values()) + sum(d_data.values())

platform_info = {
    'Sp': ('Spotify', '🟢 Spotify'),
    'Ap': ('Apple Music', '🍎 Apple Music'),
    'Yc': ('YouTube Content ID', '▶️ YouTube Official Content vs Art Tracks'),
    'Ym': ('YouTube Audio Tier', '🎵 YouTube Audio Tier'),
    'Yv': ('YouTube Music Video', '🎬 YouTube Music Video'),
    'Sh': ('YouTube Shorts', '⏱️ YouTube Shorts'),
    'Fb': ('Facebook', '👥 Facebook'),
    'Tk': ('TikTok', '🎵 TikTok'),
}

results = {}
for key, (name, emoji_name) in platform_info.items():
    b = b_data.get(key, 0)
    d = d_data.get(key, 0)
    combined = b + d
    pct_total = combined / total_rev * 100 if total_rev else 0
    pct_d = d / combined * 100 if combined else 0
    pct_b = b / combined * 100 if combined else 0
    results[key] = {
        'name': name,
        'emoji_name': emoji_name,
        'pct_total': pct_total,
        'pct_d': pct_d,
        'pct_b': pct_b,
        'combined': combined,
        'd': d,
        'b': b,
    }
    print(f"{name}: {pct_total:.1f}% total | 东西 {pct_d:.1f}% | Believe {pct_b:.1f}%")

# === Update HTML ===
with open('/Users/olivia/WorkBuddy/2026-05-13-task-4/index.html', 'r', encoding='utf-8') as f:
    html = f.read()

# Platform header replacements
replacements = {
    '🟢 Spotify': results['Sp'],
    '🍎 Apple Music': results['Ap'],
    '▶️ YouTube Official Content vs Art Tracks': results['Yc'],
    '🎵 YouTube Audio Tier': results['Ym'],
    '🎬 YouTube Music Video': results['Yv'],
    '⏱️ YouTube Shorts': results['Sh'],
    '👥 Facebook': results['Fb'],
    '🎵 TikTok': results['Tk'],
}

for emoji_name, data in replacements.items():
    old = f'{emoji_name} — 每1,000 streams'
    new = f"{emoji_name} — 每1,000 streams（占总收入 {data['pct_total']:.1f}%，其中东西占{data['pct_d']:.1f}%，Believe占{data['pct_b']:.1f}%）"
    html = html.replace(old, new)

with open('/Users/olivia/WorkBuddy/2026-05-13-task-4/index.html', 'w', encoding='utf-8') as f:
    f.write(html)

print("\nHTML updated successfully!")

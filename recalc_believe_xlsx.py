import openpyxl
from collections import defaultdict

EUR_TO_USD = 1.085

wb = openpyxl.load_workbook('/Users/olivia/Desktop/2026Q1 Believe.xlsx', data_only=True)
ws = wb.active

# Find column indices
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
platform_idx = headers.index('Platform')
gross_idx = headers.index('Gross Revenue')

platform_revenue = defaultdict(float)

for row in ws.iter_rows(min_row=2, values_only=True):
    platform = row[platform_idx]
    gross = row[gross_idx]
    if platform and gross is not None:
        try:
            platform_revenue[platform] += float(gross)
        except (ValueError, TypeError):
            pass

print("Believe Platform Revenue (EUR):")
for p, rev in sorted(platform_revenue.items(), key=lambda x: -x[1]):
    usd = rev * EUR_TO_USD
    print(f"  {p}: EUR {rev:,.2f} -> USD {usd:,.2f}")

# Map to our DATA keys
mapping = {
    'Spotify': 'Sp',
    'Apple Music': 'Ap',
    'YouTube Official Content': 'Yc',
    'YouTube UGC': 'Yc',  # combined with Official Content
    'YouTube Audio Tier': 'Ym',
    'YouTube Music Video': 'Yv',
    'YouTube Shorts': 'Sh',
    'Facebook': 'Fb',
    'TikTok': 'Tk',
}

data = {}
for platform, key in mapping.items():
    rev = platform_revenue.get(platform, 0.0)
    usd = rev * EUR_TO_USD
    if key in data:
        data[key] += usd
    else:
        data[key] = usd

print("\nBelieve DATA mapping (USD):")
for k, v in sorted(data.items()):
    print(f"  {k}: {v:,.2f}")
